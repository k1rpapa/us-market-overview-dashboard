"""米国市場俯瞰ダッシュボード用の指標データ(macro.json)を生成する。

- 取得できた値のみを出力する。取得不能・未接続の指標は値を持たず status と reason を明示する。
- EIA_API_KEY は環境変数からのみ読み込み、出力ファイルやログには含めない。
"""
import csv
import argparse
import calendar
import io
import json
import math
import os
import re
import urllib.parse
import urllib.request
from datetime import datetime, timezone

try:
    import yfinance as yf
    YFINANCE_AVAILABLE = True
except ImportError:
    YFINANCE_AVAILABLE = False

MAX_HISTORY_POINTS = 360
USER_AGENT = "Mozilla/5.0 (CapitalFlowAnalyzer macro fetch)"

GROUPS = [
    {"id": "valuation", "name": "割高感・利益"},
    {"id": "momentum", "name": "モメンタム・市場の広がり"},
    {"id": "credit", "name": "信用・センチメント"},
    {"id": "economy", "name": "景気"},
    {"id": "rates", "name": "金利"},
    {"id": "energy", "name": "エネルギー"},
]


def _spec(id, group, name, definition, reading, caution, source, source_url, frequency,
          unit="", kind="pending", params=None, reason="", formula="", restricted=False):
    return {
        "id": id, "group": group, "name": name, "definition": definition,
        "reading": reading, "caution": caution, "source": source,
        "source_url": source_url, "frequency": frequency, "unit": unit,
        "kind": kind, "params": params or {}, "reason": reason, "formula": formula,
        "restricted": restricted,
    }


FRED_URL = "https://fred.stlouisfed.org/series/"

SPECS = [
    _spec("cape", "valuation", "Shiller CAPE",
          "S&P500の株価を過去10年の平均インフレ調整後利益で割った景気循環調整PER。",
          "歴史的レンジの上位ほど長期の割高感が強い局面。",
          "短期の売買タイミング指標ではない。ソース更新停止時は最新観測値が古くなります。",
          "Robert Shiller (shillerdata.com / Yale)", "https://shillerdata.com/", "月次",
          kind="shiller"),
    _spec("erp", "valuation", "株式リスクプレミアム (ERP)",
          "株式の期待リターンが無リスク金利をどれだけ上回るか。",
          "低いほど株式の上乗せ報酬が薄い。",
          "CAPEを使った長期的な簡易代理値で、将来リターン予測や一般的なフォワードERPとは異なる。",
          "FRED (DFII10) + CAPE", FRED_URL + "DFII10", "日次/月次", unit="%",
          formula="ERP代理値(%) = 100 / Shiller CAPE − 実質10年金利 DFII10(%)。CAPE観測月内の最終実質金利を使用。",
          kind="derived", params={"left": "cape", "right": "real_yield", "method": "cape_erp"},
          reason="CAPE・実質金利が取得できない場合は算出されません。"),
    _spec("eps_revision", "valuation", "S&P500 利益予想修正幅",
          "アナリストの1株利益予想が上方/下方どちらへ修正されているかの比率。",
          "上方修正優勢なら利益モメンタムが改善。",
          "予想は楽観に偏りやすい。",
          "LSEG I/B/E/S, FactSet (商用)", "https://insight.factset.com/topic/earnings", "週次",
          reason="商用データで公的な取得手段がないため未接続。"),
    _spec("sp500_200d", "momentum", "S&P500 200日線乖離 (モメンタム)",
          "S&P500終値の200日単純移動平均からの乖離率。",
          "プラスは長期上昇トレンド、マイナスは下降トレンド。大きなプラスは過熱の目安。",
          "トレンド追随指標で転換は遅れる。", "Yahoo Finance (^GSPC)",
          "https://finance.yahoo.com/quote/%5EGSPC", "日次", "%", kind="yf_momentum", restricted=True,
          params={"ticker": "^GSPC", "window": 200}),
    _spec("ad_line", "momentum", "騰落(A/D)ライン",
          "値上がり銘柄数−値下がり銘柄数の累積。",
          "指数が上昇してもA/Dが伴わなければ上昇の裾野が狭い。",
          "取引所ごとに定義が異なる。", "NYSE / Nasdaq (商用ベンダー)",
          "https://www.nyse.com/market-data", "日次",
          reason="公的な全銘柄騰落データの安定取得手段がないため未接続。"),
    _spec("pct_above_200d", "momentum", "200日線上銘柄比率",
          "指数構成銘柄のうち終値が200日線より上にある割合。",
          "高いほど上昇の裾野が広い。極端な低下は広範な弱さ。",
          "構成銘柄の履歴データが必要で、無料APIでは再現性が低い。",
          "S&P Dow Jones Indices (商用)", "https://www.spglobal.com/spdji/en/", "日次",
          reason="構成銘柄の全履歴が必要なため未接続。"),
    _spec("new_hi_lo", "momentum", "新高値−新安値 / McClellan",
          "52週新高値銘柄数と新安値銘柄数の差、およびA/D差の19/39日EMA差(McClellan)。",
          "新高値が減る中での指数上昇はダイバージェンス。",
          "A/Dデータ依存。", "NYSE (商用ベンダー)", "https://www.nyse.com/market-data", "日次",
          reason="A/Dと同じ理由で未接続。"),
    _spec("cap_vs_equal", "momentum", "時価総額加重 vs 均等加重 (集中度)",
          "S&P500 時価総額加重ETF(SPY)に対する均等加重ETF(RSP)の比の逆数 SPY/RSP。",
          "上昇は少数の大型株への集中が進んでいることを示す。",
          "ETFの価格比であり、構成比そのものではない。配当/経費の差も含む。",
          "Yahoo Finance (SPY, RSP)", "https://finance.yahoo.com/quote/RSP", "日次", "比",
          kind="yf_ratio", restricted=True, params={"numerator": "SPY", "denominator": "RSP"},
          formula="SPY終値 / RSP終値"),
    _spec("hy_oas", "credit", "ハイイールド社債 OAS", "米国HY債の対国債オプション調整後スプレッド。",
          "拡大は信用不安・リスクオフ、極端な縮小は信用面の油断。",
          "スプレッドの水準は格付け構成の変化にも左右される。",
          "ICE BofA via FRED", FRED_URL + "BAMLH0A0HYM2", "日次", "%",
          kind="fred", restricted=True, params={"series": "BAMLH0A0HYM2"}),
    _spec("ig_oas", "credit", "投資適格社債 OAS", "米国IG債の対国債オプション調整後スプレッド。",
          "拡大は信用環境の悪化。HYより早く動かないことも多い。",
          "金利水準の影響を受けるため単独で判断しない。",
          "ICE BofA via FRED", FRED_URL + "BAMLC0A0CM", "日次", "%",
          kind="fred", restricted=True, params={"series": "BAMLC0A0CM"}),
    _spec("nfci", "credit", "シカゴ連銀 NFCI", "金融環境の総合指数。0が平均、プラスが平均より引き締まり。",
          "上昇は資金調達環境の悪化。", "週次で改定されることがある。",
          "Chicago Fed via FRED", FRED_URL + "NFCI", "週次", "指数",
          kind="fred", params={"series": "NFCI"}),
    _spec("margin_debt", "credit", "米国マージン残高", "証券会社の信用取引の借入残高。",
          "急増はレバレッジ過熱、急減は強制決済局面の兆候。",
          "月次で公表が遅れる。名目値なので株価水準との比較が必要。FINRAは月次数値が報告手法変更の影響を受ける場合があると注意喚起。",
          "FINRA", "https://www.finra.org/rules-guidance/key-topics/margin-accounts/margin-statistics",
          "月次", unit="$M", kind="finra"),
    _spec("cot_sp500", "credit", "CFTC S&P500 先物レバレッジドファンドのネットポジション",
          "CFTC Traders in Financial Futures における S&P 500 Consolidated のレバレッジドファンド買建−売建を建玉残高で割った比率。",
          "正はネットロング、負はネットショート。極端な値はポジションの偏りを示す。",
          "火曜時点の建玉を金曜公表。先物のみで現物/オプションや全市場のポジションではない。",
          "CFTC Public Reporting Environment", "https://publicreporting.cftc.gov/", "週次", "% OI",
          kind="cftc", params={"market": "S&P 500 Consolidated"}),
    _spec("fear_greed", "credit", "CNN Fear & Greed",
          "7つの市場指標を合成したセンチメント指数。",
          "極端な恐怖/強欲は逆張りの材料にされることがある。",
          "構成要素に本ダッシュボードの指標(モメンタム・VIX・Put/Call等)と重複する。公式APIは非公開。",
          "CNN Business", "https://www.cnn.com/markets/fear-and-greed", "日次", kind="link",
          reason="公式の取得APIが確認できないため未接続。リンク先で確認。"),
    _spec("put_call", "credit", "プット/コール比率", "オプション取引のプット出来高÷コール出来高。",
          "高いほど弱気/ヘッジ需要が強い。極端値は逆張りの目安。",
          "ヘッジ目的の取引も含まれる。Cboe配布ファイルはCboe Webサイト利用規約に従うためローカル版のみ。",
          "Cboe", "https://www.cboe.com/us/options/market_statistics/daily/", "日次",
          unit="比", kind="cboe_put_call", restricted=True),
    _spec("vix_curve", "credit", "VIX先物カーブ (近似: VIX/VIX3M)",
          "30日物VIXと3か月物VIX3Mの比。",
          "1超(バックワーデーション)は短期の警戒が強い。1未満は通常のコンタンゴ。",
          "先物カーブそのものではなく指数ベースの近似。",
          "Yahoo Finance (^VIX, ^VIX3M)", "https://finance.yahoo.com/quote/%5EVIX", "日次", "比",
          kind="yf_ratio", restricted=True, params={"numerator": "^VIX", "denominator": "^VIX3M"},
          formula="VIX / VIX3M"),
    _spec("yield_curve", "economy", "イールドカーブ (10年−3か月)", "米国債10年と3か月の利回り差。",
          "マイナス(逆イールド)は過去に景気後退に先行しやすい。",
          "逆イールド解消後に景気後退が来る例も多く、時期は不定。",
          "Fed via FRED", FRED_URL + "T10Y3M", "日次", "%",
          kind="fred", params={"series": "T10Y3M"}),
    _spec("recession_prob", "economy", "NY連銀 後退確率", "イールドカーブから推計した12か月先の景気後退確率。",
          "高いほど逆イールド起点の後退リスクが高いとモデルが示す。",
          "単一モデルによる推計。10年−3か月スプレッドを使い、景気後退の時期を確定する予測ではない。",
          "NY Fed", "https://www.newyorkfed.org/research/capital_markets/ycfaq.html", "月次", "%",
          kind="nyfed_recession"),
    _spec("card_delinq", "economy", "クレジットカード 30日以上延滞へのフロー",
          "NY連銀/Equifax Household Debt and Credit Report の新規30日以上延滞残高の割合。",
          "上昇は家計の返済余力低下や延滞の広がりを示す。",
          "延滞への残高フローで、カード利用者個人の延滞確率ではない。",
          "NY Fed Consumer Credit Panel/Equifax",
          "https://www.newyorkfed.org/microeconomics/databank.html", "四半期", "%",
          kind="nyfed_hhdc", params={"sheet": "Page 13 Data", "column": "CC"}),
    _spec("auto_delinq", "economy", "自動車ローン 30日以上延滞へのフロー",
          "NY連銀/Equifax Household Debt and Credit Report の新規30日以上延滞残高の割合。",
          "上昇は家計の返済余力低下や延滞の広がりを示す。",
          "延滞への残高フローで、借り手個人の延滞確率ではない。",
          "NY Fed Consumer Credit Panel/Equifax",
          "https://www.newyorkfed.org/microeconomics/databank.html", "四半期", "%",
          kind="nyfed_hhdc", params={"sheet": "Page 13 Data", "column": "AUTO"}),
    _spec("card_serious_delinq", "economy", "クレジットカード 90日以上延滞へのフロー",
          "NY連銀/Equifax Household Debt and Credit Report の新規90日以上延滞残高の割合。",
          "上昇は深刻な家計信用ストレスの拡大を示す。",
          "延滞への残高フローで、借り手個人の延滞確率ではない。",
          "NY Fed Consumer Credit Panel/Equifax",
          "https://www.newyorkfed.org/microeconomics/databank.html", "四半期", "%",
          kind="nyfed_hhdc", params={"sheet": "Page 14 Data", "column": "CC"}),
    _spec("auto_serious_delinq", "economy", "自動車ローン 90日以上延滞へのフロー",
          "NY連銀/Equifax Household Debt and Credit Report の新規90日以上延滞残高の割合。",
          "上昇は深刻な家計信用ストレスの拡大を示す。",
          "延滞への残高フローで、借り手個人の延滞確率ではない。",
          "NY Fed Consumer Credit Panel/Equifax",
          "https://www.newyorkfed.org/microeconomics/databank.html", "四半期", "%",
          kind="nyfed_hhdc", params={"sheet": "Page 14 Data", "column": "AUTO"}),
    _spec("sloos", "economy", "SLOOS 消費者貸出基準 (カード)",
          "銀行がカードローン基準を厳格化したと答えた割合−緩和した割合(ネット%)。",
          "上昇は信用供給の絞り込み。", "四半期の調査で結果は遅れて公表される。",
          "FRB SLOOS via FRED", FRED_URL + "DRTSCLCC", "四半期", "%",
          kind="fred", params={"series": "DRTSCLCC"}),
    _spec("core_capex", "economy", "コア資本財受注", "非国防資本財(航空機除く)の新規受注。設備投資の先行指標。",
          "減速は企業の投資意欲の低下。", "名目値で月次変動が大きい。",
          "US Census Bureau M3 via FRED", FRED_URL + "NEWORDER", "月次", "百万$",
          kind="fred", params={"series": "NEWORDER"}),
    _spec("real_yield", "rates", "実質金利 (10年TIPS)", "10年物インフレ連動国債の利回り。",
          "上昇は株式の割引率上昇で割高株に逆風。", "TIPSの需給でも変動する。",
          "Fed via FRED", FRED_URL + "DFII10", "日次", "%",
          kind="fred", params={"series": "DFII10"}),
    _spec("term_premium", "rates", "NY連銀 ACMタームプレミアム",
          "長期債を保有する見返りとして要求される上乗せ利回り(10年)。",
          "上昇は長期金利の上昇が期待ではなく需給/不確実性起因であることを示す。",
          "モデル推計値で改定される。公開ファイルのTERMYld列は10年ACMタームプレミアム。",
          "NY Fed", "https://www.newyorkfed.org/research/data_indicators/term-premia-tabs", "月次", "%",
          kind="acm"),
    _spec("wti", "energy", "WTI原油", "WTI原油スポット価格。",
          "急騰はインフレ・コスト圧力、急落は需要懸念の可能性。", "地政学や在庫要因で大きく振れる。",
          "EIA via FRED", FRED_URL + "DCOILWTICO", "日次", "$/bbl",
          kind="fred", params={"series": "DCOILWTICO"}),
    _spec("diesel_price", "energy", "ディーゼル小売価格", "米国平均の小売ディーゼル(No.2)価格。",
          "物流・産業コストの先行的な圧力。", "小売価格は税・流通マージンを含む。",
          "EIA via FRED", FRED_URL + "GASDESW", "週次", "$/gal",
          kind="fred", params={"series": "GASDESW"}),
    _spec("distillate_stocks", "energy", "ディスティレート在庫", "米国の留出油(ディーゼル等)在庫。",
          "低水準はディーゼル価格の上振れ要因。", "季節性が強いため同時期比較が必要。",
          "EIA API", "https://www.eia.gov/opendata/", "週次", "千bbl", kind="eia",
          params={"route": "petroleum/stoc/wstk", "series": "WDISTUS1"},
          reason="EIA_API_KEY が未設定のため取得していない。"),
    _spec("diesel_crude_spread", "energy", "ディーゼル−原油 価格差",
          "小売ディーゼルを$/bblに換算した額とWTIの差。精製マージンの粗い目安。",
          "拡大は精製・流通マージンの上昇またはディーゼル需給の逼迫。",
          "小売価格ベースで税・流通費を含み、純粋なクラックスプレッドではない。",
          "EIA via FRED", FRED_URL + "GASDESW", "週次", "$/bbl", kind="derived",
          params={"left": "diesel_price", "right": "wti", "factor": 42},
          formula="ディーゼル小売($/gal) × 42 − WTI($/bbl)"),
]


def _http_get(url, timeout=30):
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=timeout) as res:
        return res.read().decode("utf-8")


def _http_bytes(url, timeout=90):
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=timeout) as res:
        return res.read()


SHILLER_PAGE_URL = "https://shillerdata.com/"
YALE_SHILLER_URL = "http://www.econ.yale.edu/~shiller/data/ie_data.xls"


def parse_shiller_cape(content):
    import xlrd

    sheet = xlrd.open_workbook(file_contents=content).sheet_by_name("Data")
    headers = [str(v).strip() for v in sheet.row_values(7)]
    cape_col = headers.index("CAPE")
    result = []
    for row in range(8, sheet.nrows):
        raw_date = sheet.cell_value(row, 0)
        raw_value = sheet.cell_value(row, cape_col)
        try:
            date_num = float(raw_date)
            year = int(date_num)
            month = int(round((date_num - year) * 100))
            value = float(raw_value)
            if not 1 <= month <= 12 or value <= 0:
                continue
            result.append((f"{year:04d}-{month:02d}-01", value))
        except (TypeError, ValueError, OverflowError):
            continue
    if not result:
        raise ValueError("Shiller workbook has no CAPE observations")
    return result


def shiller_workbook_url(page_html):
    match = re.search(r'(?:https:)?//img1\.wsimg\.com/[^"\'\s<>]*?ie_data\.xls[^"\'\s<>]*', page_html)
    if not match:
        raise ValueError("Shiller data page has no ie_data.xls link")
    url = match.group(0).replace("&amp;", "&")
    return "https:" + url if url.startswith("//") else url


def fetch_shiller(http_get_bytes=_http_bytes, http_get=_http_get):
    """Prefer Shiller's maintained data site; the Yale copy has stopped updating."""
    try:
        return parse_shiller_cape(http_get_bytes(shiller_workbook_url(http_get(SHILLER_PAGE_URL))))
    except Exception:
        return parse_shiller_cape(http_get_bytes(YALE_SHILLER_URL))


def fetch_finra_margin(http_get_bytes=_http_bytes):
    from openpyxl import load_workbook

    content = http_get_bytes("https://www.finra.org/sites/default/files/2021-03/margin-statistics.xlsx")
    sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
    result = []
    for row in sheet.iter_rows(values_only=True):
        if len(row) < 2 or not row[0]:
            continue
        raw_date, raw_value = row[0], row[1]
        if not isinstance(raw_date, str) or not raw_date[:4].isdigit():
            continue
        try:
            year, month = raw_date.split("-")[:2]
            result.append((f"{int(year):04d}-{int(month):02d}-01", float(raw_value)))
        except (ValueError, TypeError):
            continue
    if not result:
        raise ValueError("FINRA workbook has no margin balance observations")
    return sorted(result)


def fetch_acm(http_get=_http_get):
    url = "https://www.newyorkfed.org/medialibrary/media/research/data_indicators/acmPlot_data.csv"
    reader = csv.DictReader(io.StringIO(http_get(url, timeout=60)))
    result = []
    for row in reader:
        try:
            date = datetime.strptime(row["RunDates"], "%d-%b-%Y").strftime("%Y-%m-%d")
            result.append((date, float(row["TERMYld"])))
        except (KeyError, ValueError, TypeError):
            continue
    if not result:
        raise ValueError("NY Fed ACM CSV has no term premium observations")
    return result


def fetch_nyfed_recession(http_get=_http_get):
    url = "https://www.newyorkfed.org/medialibrary/media/research/capital_markets/yield/assets/data/yield.csv"
    result = []
    for row in csv.DictReader(io.StringIO(http_get(url, timeout=60))):
        try:
            day, month, short_year = row["Date"].split("-")
            year_number = int(short_year)
            year = 1900 + year_number if year_number >= 60 else 2000 + year_number
            date = datetime.strptime(f"{day}-{month}-{year}", "%d-%b-%Y").strftime("%Y-%m-%d")
            value = float(row["Rec_prob"].strip().rstrip("%"))
            result.append((date, value))
        except (KeyError, ValueError, TypeError):
            continue
    if not result:
        raise ValueError("NY Fed recession probability CSV has no observations")
    return result


PROBIT_INTERCEPT = -0.5333
PROBIT_SLOPE = -0.6330


def recession_probit_proxy(spread_daily):
    """Estrella-Mishkin style probit used by the NY Fed: 12-month-ahead probability from the monthly mean 10y-3m spread."""
    months = {}
    for date, value in spread_daily:
        months.setdefault(date[:7], []).append((date, value))
    result = []
    for key in sorted(months):
        points = months[key]
        spread = sum(v for _, v in points) / len(points)
        z = PROBIT_INTERCEPT + PROBIT_SLOPE * spread
        result.append((points[-1][0], round(100 * 0.5 * (1 + math.erf(z / math.sqrt(2))), 2)))
    return result


BARCHART_CSV_DEFAULT = os.path.join("local_data", "barchart_cpcs.csv")


def parse_barchart_csv(text):
    """Parse a manually downloaded Barchart historical CSV (e.g. $CPCS). Footer/blank lines are skipped."""
    rows = list(csv.reader(io.StringIO(text.lstrip("\ufeff"))))
    header = None
    for index, row in enumerate(rows):
        names = [c.strip().lower() for c in row]
        date_col = next((i for i, n in enumerate(names) if n in ("time", "date")), None)
        value_col = next((i for i, n in enumerate(names) if n in ("last", "close", "latest")), None)
        if date_col is not None and value_col is not None:
            header = (index, date_col, value_col)
            break
    if header is None:
        raise ValueError("Barchart CSV needs Time/Date and Last/Close columns")
    result = []
    for row in rows[header[0] + 1:]:
        try:
            raw_date = row[header[1]].strip()
            for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
                try:
                    date = datetime.strptime(raw_date, fmt).strftime("%Y-%m-%d")
                    break
                except ValueError:
                    date = None
            if date is None:
                continue
            result.append((date, float(row[header[2]].replace(",", ""))))
        except (IndexError, ValueError):
            continue
    if not result:
        raise ValueError("Barchart CSV has no observations")
    return sorted(set(result))


def read_barchart_cpcs(path=None):
    """Return the manually imported series, or None when no file was placed locally."""
    path = path or os.environ.get("BARCHART_CPCS_CSV") or os.path.join(
        os.path.dirname(os.path.abspath(__file__)), BARCHART_CSV_DEFAULT)
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8-sig") as handle:
        return parse_barchart_csv(handle.read())


def fetch_cboe_put_call(http_get=_http_get):
    url = "https://cdn.cboe.com/resources/options/volume_and_call_put_ratios/totalpc.csv"
    rows = csv.reader(io.StringIO(http_get(url, timeout=60)))
    header = None
    result = []
    for row in rows:
        normalized = [cell.strip() for cell in row]
        if "DATE" in normalized and "P/C Ratio" in normalized:
            header = {name: normalized.index(name) for name in ("DATE", "P/C Ratio")}
            continue
        if header:
            try:
                date = datetime.strptime(normalized[header["DATE"]], "%m/%d/%Y").strftime("%Y-%m-%d")
                value = float(normalized[header["P/C Ratio"]])
                result.append((date, value))
            except (IndexError, ValueError):
                continue
    if not result:
        raise ValueError("Cboe total put/call CSV has no observations")
    return result


def fetch_cot(params, http_get=_http_get):
    query = urllib.parse.urlencode({
        "$select": "report_date_as_yyyy_mm_dd,open_interest_all,lev_money_positions_long,lev_money_positions_short",
        "$where": f"contract_market_name='{params['market']}'",
        "$order": "report_date_as_yyyy_mm_dd ASC",
        "$limit": 50000,
    })
    url = "https://publicreporting.cftc.gov/resource/gpe5-46if.json?" + query
    rows = json.loads(http_get(url, timeout=60))
    result = []
    for row in rows:
        try:
            date = row["report_date_as_yyyy_mm_dd"][:10]
            open_interest = float(row["open_interest_all"])
            long = float(row["lev_money_positions_long"])
            short = float(row["lev_money_positions_short"])
            if open_interest:
                result.append((date, round((long - short) * 100 / open_interest, 4)))
        except (KeyError, TypeError, ValueError):
            continue
    if not result:
        raise ValueError("CFTC dataset has no matching S&P 500 observations")
    return result


def parse_nyfed_hhdc_xlsx(content, sheet_name, column):
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True)[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    column_index = None
    for row in rows:
        if row and column in [str(v).strip() if v is not None else "" for v in row]:
            column_index = [str(v).strip() if v is not None else "" for v in row].index(column)
            break
    if column_index is None:
        raise ValueError(f"NY Fed workbook missing column {column} in {sheet_name}")

    result = []
    for row in rows:
        try:
            label = str(row[0]).strip()
            match = re.fullmatch(r"(\d{2}):Q([1-4])", label)
            if not match:
                continue
            year = 2000 + int(match.group(1))
            month = int(match.group(2)) * 3
            value = float(row[column_index])
            day = calendar.monthrange(year, month)[1]
            result.append((f"{year:04d}-{month:02d}-{day:02d}", value))
        except (IndexError, TypeError, ValueError):
            continue
    if not result:
        raise ValueError(f"NY Fed workbook has no observations in {sheet_name}")
    return result


def _quarter_offset(year, quarter, offset):
    index = year * 4 + quarter - 1 - offset
    return index // 4, index % 4 + 1


def fetch_nyfed_hhdc(params, http_get_bytes=_http_bytes, today=None, content_cache=None):
    today = today or datetime.now(timezone.utc).date()
    content_cache = content_cache if content_cache is not None else {}
    quarter = (today.month - 1) // 3 + 1
    last_error = None
    for offset in range(1, 9):
        year, q = _quarter_offset(today.year, quarter, offset)
        url = ("https://www.newyorkfed.org/medialibrary/interactives/householdcredit/"
               f"data/xls/hhd_c_report_{year}q{q}.xlsx")
        try:
            if url not in content_cache:
                content_cache[url] = http_get_bytes(url)
            return parse_nyfed_hhdc_xlsx(content_cache[url], params["sheet"], params["column"])
        except Exception as exc:
            last_error = exc
    raise ValueError(f"NY Fed household credit workbook unavailable ({type(last_error).__name__})")


def parse_fred_csv(text):
    series = []
    reader = csv.reader(io.StringIO(text))
    next(reader, None)
    for row in reader:
        if len(row) < 2:
            continue
        try:
            series.append((row[0], float(row[1])))
        except ValueError:
            continue
    return series


def fetch_fred(series_id, http_get=_http_get):
    """FRED_API_KEY があれば公式API、なければ公開CSVから取得する。"""
    key = os.environ.get("FRED_API_KEY")
    if key:
        query = urllib.parse.urlencode({"series_id": series_id, "api_key": key, "file_type": "json"})
        payload = json.loads(http_get("https://api.stlouisfed.org/fred/series/observations?" + query))
        series = []
        for o in payload.get("observations", []):
            try:
                series.append((o["date"], float(o["value"])))
            except (KeyError, ValueError):
                continue
    else:
        url = "https://fred.stlouisfed.org/graph/fredgraph.csv?id=" + urllib.parse.quote(series_id)
        series = parse_fred_csv(http_get(url))
    if not series:
        raise ValueError("empty series")
    return series

def fetch_eia(params, http_get=_http_get):
    key = os.environ.get("EIA_API_KEY")
    if not key:
        raise RuntimeError("EIA_API_KEY not set")
    query = urllib.parse.urlencode([
        ("api_key", key), ("frequency", "weekly"), ("data[0]", "value"),
        ("facets[series][]", params["series"]), ("sort[0][column]", "period"),
        ("sort[0][direction]", "asc"), ("length", "5000"),
    ])
    payload = json.loads(http_get(f"https://api.eia.gov/v2/{params['route']}/data/?{query}"))
    rows = payload.get("response", {}).get("data", [])
    series = [(r["period"], float(r["value"])) for r in rows if r.get("value") is not None]
    if not series:
        raise ValueError("empty series")
    return series


def _yf_close(ticker):
    if not YFINANCE_AVAILABLE:
        raise RuntimeError("yfinance not installed")
    df = yf.download(ticker, period="max", progress=False, auto_adjust=True)["Close"]
    if hasattr(df, "columns"):
        df = df.iloc[:, 0]
    df = df.dropna()
    if len(df) == 0:
        raise ValueError("empty series")
    return [(d.strftime("%Y-%m-%d"), float(v)) for d, v in df.items()]


def moving_average_deviation(series, window):
    out = []
    total = 0.0
    for i, (d, v) in enumerate(series):
        total += v
        if i >= window:
            total -= series[i - window][1]
        if i >= window - 1:
            out.append((d, round((v / (total / window) - 1) * 100, 3)))
    return out


def ratio_series(num, den):
    den_map = dict(den)
    return [(d, round(v / den_map[d], 5)) for d, v in num if den_map.get(d)]


def cape_erp(cape, real_yield):
    """CAPE earnings-yield proxy less the last real yield observation in that month."""
    result = []
    yield_index = 0
    for date, cape_value in cape:
        year, month = map(int, date[:7].split("-"))
        next_month = datetime(year + (month == 12), month % 12 + 1, 1).strftime("%Y-%m-%d")
        while yield_index < len(real_yield) and real_yield[yield_index][0] < next_month:
            yield_index += 1
        candidate_index = yield_index - 1
        if candidate_index >= 0 and real_yield[candidate_index][0][:7] == date[:7]:
            result.append((date, round(100 / cape_value - real_yield[candidate_index][1], 4)))
    return result


def derived_spread(left, right, factor):
    """週次のleftに対し、同日以前の直近のright値を使って left*factor - right を返す。"""
    out = []
    j = -1
    for d, v in left:
        while j + 1 < len(right) and right[j + 1][0] <= d:
            j += 1
        if j >= 0:
            out.append((d, round(v * factor - right[j][1], 3)))
    return out


def downsample(series, limit=MAX_HISTORY_POINTS):
    if len(series) <= limit:
        return [list(p) for p in series]
    step = (len(series) - 1) / (limit - 1)
    idx = sorted({round(i * step) for i in range(limit)} | {len(series) - 1})
    return [list(series[i]) for i in idx]


def summarize(series):
    values = [v for _, v in series]
    latest_date, latest_value = series[-1]
    ordered = sorted(values)
    n = len(ordered)
    median = ordered[n // 2] if n % 2 else (ordered[n // 2 - 1] + ordered[n // 2]) / 2
    percentile = round(100 * sum(1 for v in values if v <= latest_value) / n, 1)
    return {
        "latest": {"date": latest_date, "value": round(latest_value, 4)},
        "stats": {"min": round(ordered[0], 4), "max": round(ordered[-1], 4),
                  "median": round(median, 4), "percentile": percentile,
                  "count": n, "start": series[0][0]},
        "history": downsample(series),
    }


def _public(spec):
    keys = ("id", "group", "name", "definition", "reading", "caution", "source",
            "source_url", "frequency", "unit", "formula")
    return {k: spec[k] for k in keys}


RESTRICTED_REASON = "第三者(指数提供元/データ配信元)の再配布制限があるため、公開リポジトリでは値・履歴を掲載しません。出典リンク先で確認してください。"
LOCAL_RESTRICTED_REASON = "ローカル専用データです。公開版への再配布はできません。"


def build_indicator(spec, series=None, error=None, allow_restricted=False):
    item = _public(spec)
    if spec.get("restricted") and not allow_restricted:
        item.update({"status": "restricted", "reason": RESTRICTED_REASON})
    elif series:
        item.update({"status": "ok", "reason": ""})
        item.update(summarize(series))
        if spec.get("restricted"):
            item["caution"] += " " + LOCAL_RESTRICTED_REASON
    elif spec["kind"] == "link":
        item.update({"status": "link_only", "reason": spec["reason"]})
    elif spec["kind"] == "pending":
        item.update({"status": "pending", "reason": spec["reason"]})
    else:
        item.update({"status": "unavailable", "reason": error or spec["reason"] or "取得に失敗しました。"})
    return item


def generate(fred=fetch_fred, eia=fetch_eia, yf_close=_yf_close, now=None, allow_restricted=False,
             shiller=fetch_shiller, finra=fetch_finra_margin, acm=fetch_acm, cot=fetch_cot,
             nyfed_hhdc=fetch_nyfed_hhdc, recession=fetch_nyfed_recession,
             cboe=fetch_cboe_put_call, barchart=read_barchart_cpcs):
    cache = {}
    results = {}
    errors = {}
    hhdc_content_cache = {}
    proxies = {}

    def fred_cached(sid):
        if sid not in cache:
            cache[sid] = fred(sid)
        return cache[sid]

    for spec in SPECS:
        kind, p = spec["kind"], spec["params"]
        if spec.get("restricted") and not allow_restricted:
            continue
        try:
            if kind == "fred":
                results[spec["id"]] = fred_cached(p["series"])
            elif kind == "eia":
                results[spec["id"]] = eia(p)
            elif kind == "yf_momentum":
                results[spec["id"]] = moving_average_deviation(yf_close(p["ticker"]), p["window"])
            elif kind == "yf_ratio":
                results[spec["id"]] = ratio_series(yf_close(p["numerator"]), yf_close(p["denominator"]))
            elif kind == "shiller":
                results[spec["id"]] = shiller()
            elif kind == "finra":
                results[spec["id"]] = finra()
            elif kind == "acm":
                results[spec["id"]] = acm()
            elif kind == "nyfed_recession":
                official = None
                try:
                    official = recession()
                except Exception:
                    pass
                limit = (now or datetime.now(timezone.utc)).date()
                if official and (limit - datetime.strptime(official[-1][0], "%Y-%m-%d").date()).days <= 180:
                    results[spec["id"]] = official
                else:
                    results[spec["id"]] = recession_probit_proxy(fred_cached("T10Y3M"))
                    proxies[spec["id"]] = True
            elif kind == "cboe_put_call":
                manual = barchart()
                if manual:
                    results[spec["id"]] = manual
                    proxies[spec["id"] + "_barchart"] = True
                else:
                    results[spec["id"]] = cboe()
            elif kind == "cftc":
                results[spec["id"]] = cot(p)
            elif kind == "nyfed_hhdc":
                results[spec["id"]] = nyfed_hhdc(p, content_cache=hhdc_content_cache)
        except Exception as exc:
            # 例外文にURLやキーが含まれ得るため、型名のみ記録する
            errors[spec["id"]] = f"取得に失敗しました ({type(exc).__name__})。"
            if kind == "eia" and "EIA_API_KEY" in str(exc):
                errors[spec["id"]] = spec["reason"]

    for spec in SPECS:
        if spec["kind"] == "derived" and (allow_restricted or not spec.get("restricted")):
            p = spec["params"]
            if p.get("method") == "cape_erp":
                if p["left"] in results and p["right"] in results:
                    results[spec["id"]] = cape_erp(results[p["left"]], results[p["right"]])
                else:
                    errors[spec["id"]] = "CAPE・実質金利の元データを取得できず算出できません。"
            elif p["left"] in results and p["right"] in results:
                results[spec["id"]] = derived_spread(results[p["left"]], results[p["right"]], p["factor"])
            else:
                errors[spec["id"]] = "元データの取得に失敗したため算出できません。"

    generated_dt = now or datetime.now(timezone.utc)
    generated = generated_dt.strftime("%Y-%m-%dT%H:%M:%SZ")
    indicators = [build_indicator(s, results.get(s["id"]), errors.get(s["id"]), allow_restricted)
                  for s in SPECS]
    for item in indicators:
        if proxies.get(item["id"]) and item.get("latest"):
            item["name"] += " (??: ??probit??)"
            item["source"] = "FRED T10Y3M (Fed H.15) ?????? / ??NY???????"
            item["source_url"] = FRED_URL + "T10Y3M"
            item["formula"] = "??(%) = 100 ? ?(?0.5333 ? 0.6330 ? 10??3??????????%)???????????????"
            item["caution"] = ("NY????CSV??????????????????????????????????"
                               "??CSV?????(?2017-06)?????????0.07????????????????"
                               "???????????????????12?????????????????? " + item["caution"])
            item["frequency"] = "?????????"
        if proxies.get(item["id"] + "_barchart") and item.get("latest"):
            item["name"] = "???/????? (Barchart $CPCS ??CSV)"
            item["definition"] = "Barchart???(?????)????? ???/???????? $CPCS??????????"
            item["source"] = "Barchart $CPCS (?????????????CSV???????)"
            item["source_url"] = "https://www.barchart.com/stocks/quotes/$CPCS/historical-download"
            item["caution"] = ("Barchart????????????????CSV???????????????????????????? "
                               + item["caution"])
            item["frequency"] = "?? (????)"
        if item.get("latest") and item["id"] in ("cape", "erp", "recession_prob", "put_call"):
            try:
                observation = datetime.strptime(item["latest"]["date"], "%Y-%m-%d").date()
                if (generated_dt.date() - observation).days > 180:
                    item["status"] = "stale"
                    item["reason"] = "取得元ファイルの最終観測が180日超前です。数値は古く、現在値として扱わないでください。"
            except ValueError:
                continue
    return {
        "generated_at": generated,
        "groups": GROUPS,
        "indicators": indicators,
    }


def load_dotenv(path):
    """Load simple KEY=VALUE entries without overriding already-set environment variables."""
    if not os.path.exists(path):
        return
    with open(path, encoding="utf-8") as env_file:
        for raw_line in env_file:
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            key = key.strip()
            value = value.strip().strip("\"'")
            if key and key not in os.environ:
                os.environ[key] = value


def main(argv=None):
    parser = argparse.ArgumentParser(description="Fetch US market overview indicators.")
    parser.add_argument("--local", action="store_true",
                        help="include restricted series and write gitignored macro.local.json")
    args = parser.parse_args(argv)
    base_dir = os.path.dirname(os.path.abspath(__file__))
    load_dotenv(os.path.join(base_dir, ".env"))
    output_name = "macro.local.json" if args.local else "macro.json"
    out_path = os.path.join(base_dir, output_name)
    data = generate(allow_restricted=args.local)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    counts = {}
    for i in data["indicators"]:
        counts[i["status"]] = counts.get(i["status"], 0) + 1
    print(f"{output_name} written:", counts)


if __name__ == "__main__":
    main()
